"""Export reviews / chats to Markdown, HTML and DOCX."""

from __future__ import annotations

import html
import io
import re


def review_to_markdown(title: str, sections: dict[str, dict]) -> str:
    out = [f"# {title}\n"]
    for sec in sections.values():
        out.append(f"## {sec['title']}\n\n{sec['content'].strip()}\n")
    return "\n".join(out)


def markdown_to_html(md: str, title: str = "ResearchMind export") -> str:
    try:
        import markdown as mdlib  # optional dependency

        body = mdlib.markdown(md, extensions=["tables"])
    except Exception:
        body = ""
        for block in md.split("\n\n"):
            b = html.escape(block.strip())
            if b.startswith("### "):
                body += f"<h3>{b[4:]}</h3>"
            elif b.startswith("## "):
                body += f"<h2>{b[3:]}</h2>"
            elif b.startswith("# "):
                body += f"<h1>{b[2:]}</h1>"
            elif b:
                body += f"<p>{b.replace(chr(10), '<br>')}</p>"
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            "<style>body{font-family:Georgia,serif;max-width:820px;margin:40px auto;line-height:1.6;padding:0 16px}"
            "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:6px}</style></head>"
            f"<body>{body}</body></html>")


def markdown_to_docx(md: str) -> bytes:
    import docx

    d = docx.Document()
    for line in md.splitlines():
        s = line.rstrip()
        if not s:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif re.match(r"^\s*[-*•]\s+", s):
            _add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", s))
        elif re.match(r"^\s*\d+\.\s+", s):
            _add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", s))
        else:
            _add_runs(d.add_paragraph(), s)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _add_runs(par, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part:
            par.add_run(part)
